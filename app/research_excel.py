"""Excel export of a completed research snapshot; no broker calls or orders."""
from datetime import datetime, timezone
from io import BytesIO
import math

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def excel_report(report):
    if not isinstance(report, dict) or not all(isinstance(report.get(k), dict) for k in ('settings', 'segments', 'baseline_holdout')):
        raise ValueError('Tamamlanmış bir araştırma raporu gerekli.')
    settings = report['settings']
    currency = settings.get('currency', '')
    book = Workbook(write_only=True)

    def sheet(name, headers, widths):
        ws = book.create_sheet(name)
        ws.freeze_panes = 'A2'
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        cells = []
        for title in headers:
            cell = WriteOnlyCell(ws, value=title)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='164E63')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cells.append(cell)
        ws.append(cells)
        ws.row_dimensions[1].height = 32
        return ws

    def row(ws, values):
        cells = []
        for value in values:
            if value is not None and not isinstance(value, (str, int, float, bool, datetime)):
                raise ValueError('Rapor hücreleri metin veya sayı olmalı.')
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError('Rapor sonlu sayılar içermeli.')
            if isinstance(value, str) and len(value) > 32000:
                raise ValueError('Rapor hücresi çok uzun.')
            cell = WriteOnlyCell(ws, value=value)
            if isinstance(value, str):
                cell.data_type = 's'  # Export strings literally, including = and + prefixes.
            elif isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm'
            elif isinstance(value, float):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
            cells.append(cell)
        ws.append(cells)

    def date(value):
        if value is None:
            return None
        return datetime.fromtimestamp(value, timezone.utc).replace(tzinfo=None)

    def config_label(config):
        return f"{config['mode']} · {'filtreli' if config['filtered'] else 'ham'} · {config['wait']} mum · %{round(config['distance_atr']*100)} ATR"

    try:
        columns = [('count','İşlem'), ('wins','Kârlı'), ('losses','Zararlı'), ('breakeven','Başa baş'),
                   ('initial_capital','Ana para'), ('gross_profit','Toplam kâr'), ('gross_loss','Toplam zarar'),
                   ('net_profit','Net kazanç/kayıp'), ('final_capital','Son bakiye'), ('return_pct','Getiri (%)'),
                   ('max_drawdown_money','Azami düşüş'), ('win_rate_pct','Kazanma (%)'), ('net_points','Net puan')]
        summary = sheet('Özet', ['Dönem','Başlangıç (UTC)','Bitiş (UTC, hariç)','Lot','Para birimi'] + [v for _,v in columns], [30,22,22,12,14]+[18]*len(columns))
        groups = []
        selected = report.get('selected')
        if selected:
            groups.extend([('Seçilen · eğitim','train',selected['train'],report.get('selected_trades',{}).get('train',[])),
                           ('Seçilen · doğrulama','validation',selected['validation'],report.get('selected_trades',{}).get('validation',[])),
                           ('Seçilen · ayrılmış test','holdout',report['selected_holdout']['stats'],report['selected_holdout']['trades'])])
        baseline = report['baseline_holdout']
        groups.append(('Ham kesişim · ayrılmış test','holdout',baseline['stats'],baseline['trades']))
        for label, period, stats, _ in groups:
            segment = report['segments'][period]
            row(summary, [label,date(segment['start']),date(segment['end_exclusive']),settings.get('lot_size'),currency] + [stats.get(key) for key,_ in columns])
        summary.auto_filter.ref = f'A1:R{len(groups)+1}'

        inputs = sheet('Ayarlar', ['Ayar / bilgi','Değer'], [45,100])
        labels = {'hma_period':'HMA periyodu','kama_period':'KAMA ER periyodu','kama_fast':'KAMA hızlı periyodu',
                  'kama_slow':'KAMA yavaş periyodu','atr_period':'ATR periyodu','er_min':'Minimum ER',
                  'lot_size':'İşlem lotu','initial_capital':'Başlangıç sermayesi','currency':'Para birimi',
                  'point':'Broker point','contract_size':'Sözleşme büyüklüğü','stop_atr':'Stop ATR çarpanı',
                  'target_atr':'Hedef ATR çarpanı','historical_spread':'Tarihsel spread',
                  'commission_points':'Komisyon (puan/işlem)','slippage_points':'Kayma (puan/gerçekleşme)',
                  'swap_long_points_per_day':'Alış swap maliyeti (puan/gün)',
                  'swap_short_points_per_day':'Satış swap maliyeti (puan/gün)','min_train_trades':'Minimum eğitim işlemi'}
        if len(settings) > 100:
            raise ValueError('Çok fazla ayar.')
        for key, value in settings.items():
            row(inputs, [labels.get(key,key),value])
        row(inputs,['Sembol',report.get('symbol','XAUUSD')])
        row(inputs,['Seçilen aday',config_label(selected['config']) if selected else 'Uygun aday yok'])
        row(inputs,['Değerlendirme',report.get('assessment',{}).get('message')])
        row(inputs,['Kaynak','MT5 broker geçmişinden hesaplanan araştırma raporu; mevcut raporun Excel kopyası.'])
        row(inputs,['Dönemler','İlk %60 eğitim, sonraki %20 doğrulama, son %20 ayrılmış test. Her dönem aynı sermayeyle başlar.'])
        row(inputs,['İşlem sınırı','10 işlem sınırı yoktur. İşlem sayısı kesişimlere, filtrelere ve veri dönemine bağlıdır.'])
        row(inputs,['Varsayım',report.get('assumptions',{}).get('risk_note')])

        trades_sheet = sheet('İşlemler', ['Dönem','#','Yön','Lot','Giriş (UTC)','Çıkış (UTC)','Giriş fiyatı','Çıkış fiyatı',
            'Net puan','Net kazanç/kayıp','İşlem sonrası bakiye','Para birimi','Sonuç','Çıkış nedeni'], [30,10,12,12,22,22,18,18,18,20,24,14,14,24])
        total_rows = 0
        reasons = {'stop':'Stop','target':'Hedef','opposite crossover':'Ters kesişim','segment end':'Dönem sonu'}
        for label, _, stats, trades in groups:
            if not isinstance(trades,list) or len(trades) != stats['count']:
                raise ValueError('İşlem dökümü işlem sayısıyla eşleşmiyor; raporu yeniden çalıştırın.')
            total_rows += len(trades)
            if total_rows > 100000:
                raise ValueError('Excel için en fazla 100000 işlem desteklenir.')
            for i, trade in enumerate(trades,1):
                net = trade['net_points']
                row(trades_sheet,[label,i,trade['side'],trade.get('lot_size'),date(trade['entry_time']),date(trade['exit_time']),
                    trade.get('entry_price'),trade.get('exit_price'),net,trade.get('net_profit'),trade.get('balance_after'),
                    currency,'Kârlı' if net > 0 else 'Zararlı' if net < 0 else 'Başa baş',reasons.get(trade['reason'],trade['reason'])])
        trades_sheet.auto_filter.ref = f'A1:N{total_rows+1}'
        comparison = sheet('Karşılaştırma', ['Seçenek','Dönem']+[v for _,v in columns]+['Kesişim','Onay','Filtre iptali','Süre aşımı','Ters iptal'], [45,18]+[18]*len(columns)+[18]*5)
        comparisons = report.get('comparisons',[])
        if not isinstance(comparisons,list) or len(comparisons) > 100:
            raise ValueError('Geçersiz karşılaştırma sayısı.')
        for item in comparisons:
            for period, name in [('train','Eğitim'),('validation','Doğrulama')]:
                signals = item.get(period+'_signals',{})
                row(comparison,[config_label(item['config']),name]+[item[period].get(key) for key,_ in columns]+[
                    signals.get(k) for k in ('crossovers','confirmed','filtered_at_deadline','expired','reversed')])
        comparison.auto_filter.ref = f'A1:T{2*len(comparisons)+1}'
    except Exception:
        book.save(BytesIO())
        raise
    output = BytesIO()
    book.save(output)
    return output.getvalue()
