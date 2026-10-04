from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.research_excel import excel_report, MIME
from app.strategy_research import compare, Settings
from tests.test_strategy_research import history, aggregate


def report():
    rows = history()
    return compare(rows, aggregate(rows), Settings(min_train_trades=1, contract_size=100, currency='USD'))


def test_workbook_keeps_numeric_stats_and_every_trade():
    data = report()
    book = load_workbook(BytesIO(excel_report(data)))
    assert book.sheetnames == ['Özet','Ayarlar','İşlemler','Karşılaştırma']
    groups = [data['selected']['train'], data['selected']['validation'], data['selected_holdout']['stats'], data['baseline_holdout']['stats']]
    assert book['İşlemler'].max_row == 1 + sum(s['count'] for s in groups)
    assert book['Özet']['N4'].value == data['selected_holdout']['stats']['final_capital']
    assert book['Özet']['F4'].value == data['selected_holdout']['stats']['count']
    assert book['Özet']['L4'].data_type == 'n'
    assert book['İşlemler']['E2'].data_type == 'd'
    assert book['Karşılaştırma'].max_row == 103
    assert book['İşlemler'].freeze_panes == 'A2'
    assert book['İşlemler'].auto_filter.ref.endswith(str(book['İşlemler'].max_row))


def test_export_strings_never_become_excel_formulas():
    data=report();data['assessment']['message']='=HYPERLINK("example","click")'
    book=load_workbook(BytesIO(excel_report(data)))
    cells=list(book['Ayarlar'].iter_rows())
    cell=next(row[1] for row in cells if row[0].value=='Değerlendirme')
    assert cell.value.startswith('=HYPERLINK') and cell.data_type=='s'


def test_export_api_auth_validation_and_viewer_access(monkeypatch):
    import app.main as main
    web=TestClient(main.app);data=report()
    assert web.post('/api/research/export',json=data).status_code==401
    auth=('test','test-panel-password')
    response=web.post('/api/research/export',json=data,auth=auth)
    assert response.status_code==200
    assert response.headers['content-type']==MIME
    assert '.xlsx' in response.headers['content-disposition']
    assert load_workbook(BytesIO(response.content))['Özet'].max_row==5
    assert web.post('/api/research/export',json={},auth=auth).status_code==422
    monkeypatch.setenv('DASHBOARD_USER_2','reader');monkeypatch.setenv('DASHBOARD_PASSWORD_2','reader-pass');monkeypatch.setenv('DASHBOARD_ROLE_2','VIEWER')
    assert web.post('/api/research/export',json=data,auth=('reader','reader-pass')).status_code==200


def test_margin_report_exports_rules_risk_counts_and_stopout_reason():
    from dataclasses import replace
    from tests.test_strategy_research import margin_scenario
    from app.strategy_research import run_segment
    rows, study, settings = margin_scenario()
    rows[3]['low'] = 50
    segment = run_segment(rows, study, 3600, 0, 36000, settings)
    data = {'settings': __import__('dataclasses').asdict(settings),
            'segments': {'holdout': {'start':0,'end_exclusive':36000}},
            'baseline_holdout': segment, 'comparisons': []}
    book = load_workbook(BytesIO(excel_report(data)))
    headers = [cell.value for cell in book['Özet'][1]]
    assert book['Özet'].cell(2,headers.index('Stop-out')+1).value == 1
    assert book['İşlemler']['N2'].value == 'Stop-out'
    rules = next(row[1].value for row in book['Ayarlar'].iter_rows() if row[0].value == 'broker_rules')
    assert __import__('json').loads(rules)['stop_out'] == 20
